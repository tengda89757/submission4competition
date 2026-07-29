"""
read_sop_docx.py — Extract the AUTHORITATIVE task definitions from the official
case SOP .docx files (official ruling: the docx overrides task_config/sop*.md),
plus each scene's input/output port centers from its own semantic map.

Usage:  .venv\\Scripts\\python pipeline\\read_sop_docx.py
"""
from __future__ import annotations

import sys
from pathlib import Path

APP_DIR = Path(__file__).resolve().parents[1]
REPO = APP_DIR.parent
for _p in (APP_DIR / "src", APP_DIR):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

from docx import Document  # python-docx


def dump_docx(path: Path) -> None:
    print(f"\n{'='*72}\n{path.name}\n{'='*72}")
    doc = Document(str(path))
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            print(t)
    for ti, table in enumerate(doc.tables):
        print(f"--- table {ti} ---")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            print(" | ".join(cells))


def dump_ports() -> None:
    from robot_agent.core.map_loader import load_map_files
    from robot_agent.core.scene_context import SceneContext
    map_dir = APP_DIR / "robosuite" / "robosuite" / "environments" / "factory_sorting" / "generated_maps"
    prefixes = [
        "factory_sorting_1_3fo3erfhisem", "factory_sorting_3_3fo3errph7x9",
        "factory_sorting_5_3fo3ertpxeut", "factory_sorting_7_3fo3erfky9rn",
        "factory_sorting_9_3fo3ert2c5fp",
    ]
    print(f"\n{'='*72}\nSCENE PORT CENTERS (from each scene's own semantic map)\n{'='*72}")
    for pf in prefixes:
        sem = map_dir / f"{pf}_scene_regenerated_semantic_map.json"
        grid = map_dir / f"{pf}_scene_regenerated_occupancy_grid.npy"
        if not sem.exists():
            print(f"{pf}: MISSING MAP")
            continue
        scene, _ = load_map_files(sem, grid)
        ctx = SceneContext.from_semantic_map(scene)
        ins = {k: tuple(round(float(x), 2) for x in v.center[:2]) for k, v in ctx.input_ports.items()}
        outs = {k: tuple(round(float(x), 2) for x in v.center[:2]) for k, v in ctx.output_ports.items()}
        print(f"{pf}\n  in : {ins}\n  out: {outs}")


def main() -> int:
    sop_dir = REPO / "competition description" / "sop+prompt"
    for case in (1, 3, 5, 7, 9):
        p = sop_dir / f"JCIIOT 2026 case {case} SOP.docx"
        if p.exists():
            dump_docx(p)
        else:
            print(f"missing: {p}")
    dump_ports()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
